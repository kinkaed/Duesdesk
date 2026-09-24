import io
import json
from datetime import date
from decimal import Decimal
from uuid import uuid4
from django.test import TestCase, override_settings
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core import mail
from openpyxl import load_workbook
from .models import Member, UserAccess, Payment, AuditEvent, ImportBatch
from .services import record_payment, void_payment
from .views import member_rows

@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'], EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend', STORAGES={'default':{'BACKEND':'django.core.files.storage.FileSystemStorage'},'staticfiles':{'BACKEND':'django.contrib.staticfiles.storage.StaticFilesStorage'}})
class OperationalTests(TestCase):
    def setUp(self):
        self.secretary=User.objects.create_user('sec',password='A-Fresh-Strong-Password!',is_staff=True,email='sec@example.com')
        self.member=Member.objects.create(full_name='Sample',joined=date(2026,9,1))
        self.other=Member.objects.create(full_name='Other',joined=date(2026,9,1))
        self.reader=User.objects.create_user('member',password='A-Fresh-Strong-Password!',email='member@example.com')
        UserAccess.objects.create(user=self.reader,role='member',member=self.member)
        self.auditor=User.objects.create_user('auditor')
        UserAccess.objects.create(user=self.auditor,role='auditor')
        self.payload={'member_id':self.member.pk,'amount':'100','start_month':'2026-09','payment_date':'2026-09-01','method':'Cash','request_key':str(uuid4())}
        self.client.force_login(self.secretary)

    def post(self,url,data):return self.client.post(url,json.dumps(data),content_type='application/json')

    def test_member_isolation_on_all_read_paths(self):
        own=record_payment(self.payload,self.secretary)
        other=record_payment({**self.payload,'member_id':self.other.pk,'request_key':str(uuid4())},self.secretary)
        self.client.force_login(self.reader)
        self.assertEqual(len(self.client.get('/api/overview/?month=2026-09').json()['members']),1)
        self.assertEqual(self.client.get(f'/api/members/{self.other.pk}/').status_code,404)
        self.assertEqual(self.client.get(f'/receipts/{other.pk}/').status_code,404)
        self.assertEqual(self.client.get(f'/receipts/{own.pk}/').status_code,200)
        self.assertEqual(len(self.client.get('/api/payments/').json()['payments']),1)
        self.assertNotIn('Other',self.client.get('/export/?month=2026-09').content.decode())
        self.assertEqual(self.post('/api/payments/',self.payload).status_code,403)
        self.assertEqual(self.client.get('/api/accounts/').status_code,403)
        self.assertEqual(self.client.get('/api/audit/').status_code,403)

    def test_auditor_read_only(self):
        self.client.force_login(self.auditor)
        self.assertEqual(len(self.client.get('/api/members/').json()['members']),2)
        self.assertEqual(self.post('/api/members/',{'name':'No'}).status_code,403)
        self.assertEqual(self.client.get('/api/audit/').status_code,200)
        self.assertEqual(self.client.get('/api/settings/').status_code,403)

    def test_void_preserves_receipt_restores_balance_and_audit(self):
        p=record_payment(self.payload,self.secretary)
        void_payment(p.pk,self.secretary,'Wrong covered period')
        self.assertEqual(Payment.objects.count(),1)
        row=member_rows(date(2026,9,1))[1] if self.member.full_name=='Z' else next(x for x in member_rows(date(2026,9,1)) if x['id']==self.member.pk)
        self.assertEqual(Decimal(row['balance']),25)
        self.assertEqual(AuditEvent.objects.filter(action='payment.voided').count(),1)
        self.assertContains(self.client.get(f'/receipts/{p.pk}/'),'VOID')
        with self.assertRaises(ValueError):void_payment(p.pk,self.secretary,'Duplicate void')

    def test_changed_duplicate_request_is_rejected(self):
        record_payment(self.payload,self.secretary)
        with self.assertRaises(ValueError):record_payment({**self.payload,'start_month':'2026-10'},self.secretary)

    def test_receipt_name_survives_member_rename(self):
        p=record_payment(self.payload,self.secretary)
        self.member.full_name='Changed Name';self.member.save()
        self.assertContains(self.client.get(f'/receipts/{p.pk}/'),'Sample')

    def test_ending_membership_preserves_old_arrears(self):
        self.member.billing_end=date(2026,10,1);self.member.status='Inactive';self.member.save()
        row=next(x for x in member_rows(date(2026,12,1)) if x['id']==self.member.pk)
        self.assertEqual(Decimal(row['arrears']),50)
        self.assertEqual(row['status'],'Not due')

    def test_edit_member_and_audit(self):
        data={'name':'Updated','phone':'024','email':'','joined':'2026-09-01','status':'Active'}
        self.assertEqual(self.post(f'/api/members/{self.member.pk}/',data).status_code,200)
        self.assertTrue(AuditEvent.objects.filter(action='member.updated').exists())

    def test_excel_is_scoped_and_formula_safe(self):
        self.member.full_name='=SUM(1,2)';self.member.save()
        self.client.force_login(self.reader)
        response=self.client.get('/export/excel/?month=2026-09')
        book=load_workbook(io.BytesIO(response.content));sheet=book.active
        self.assertEqual(sheet.max_row,2)
        self.assertEqual(sheet['B2'].data_type,'s')
        self.assertTrue(sheet['B2'].value.startswith("'="))

    def test_import_preview_commit_and_duplicate_prevention(self):
        file=SimpleUploadedFile('members.csv',b'name,joined\nImported Person,2026-09-01\n')
        response=self.client.post('/api/import/preview/',{'kind':'members','file':file})
        self.assertEqual(response.status_code,200,response.content)
        token=response.json()['token']
        self.assertEqual(Member.objects.count(),2)
        self.assertEqual(self.post('/api/import/commit/',{'token':token}).status_code,200)
        self.assertEqual(Member.objects.count(),3)
        self.assertEqual(self.post('/api/import/commit/',{'token':token}).status_code,400)
        self.assertEqual(ImportBatch.objects.count(),1)

    def test_import_invalid_row_does_not_write(self):
        file=SimpleUploadedFile('members.csv',b'name,joined\nGood,2026-09-01\nBad,not-a-date\n')
        self.assertEqual(self.client.post('/api/import/preview/',{'kind':'members','file':file}).status_code,400)
        self.assertEqual(Member.objects.count(),2)

    def test_import_payments_atomic_and_recalculates(self):
        file=SimpleUploadedFile('payments.csv',f'member_id,amount,start_month,payment_date,method\n{self.member.pk},10,2026-09,2026-09-01,Cash\n{self.member.pk},15,2026-09,2026-09-01,Cash\n'.encode())
        token=self.client.post('/api/import/preview/',{'kind':'payments','file':file}).json()['token']
        self.assertEqual(self.post('/api/import/commit/',{'token':token}).status_code,200)
        row=next(x for x in member_rows(date(2026,9,1)) if x['id']==self.member.pk)
        self.assertEqual(row['status'],'Paid')

    def test_user_creation_validates_password_and_role(self):
        data={'username':'newuser','email':'new@example.com','password':'short','role':'member','member_id':self.other.pk}
        self.assertEqual(self.post('/api/accounts/',data).status_code,400)
        data['password']='An-Excellent-Private-Phrase!'
        self.assertEqual(self.post('/api/accounts/',data).status_code,201)
        new=User.objects.get(username='newuser')
        self.assertFalse(new.is_superuser)
        self.assertEqual(new.access.member,self.other)
        self.assertNotIn(data['password'],''.join(AuditEvent.objects.values_list('details',flat=True)))

    def test_disabling_user_blocks_existing_session(self):
        self.assertEqual(self.post(f'/api/accounts/{self.reader.pk}/disable/',{}).status_code,200)
        self.client.force_login(self.reader)
        self.assertEqual(self.client.get('/api/overview/').status_code,401)
        self.assertEqual(self.post(f'/api/accounts/{self.secretary.pk}/disable/',{}).status_code,401)

    def test_password_recovery_and_throttle(self):
        self.client.logout()
        self.client.post('/account/reset/',{'email':'member@example.com'})
        self.assertEqual(len(mail.outbox),1)
        self.client.post('/account/reset/',{'email':'member@example.com'})
        self.assertEqual(len(mail.outbox),1)

    def test_login_locks_after_failed_attempts(self):
        self.client.logout()
        for _ in range(5):self.client.post('/login/',{'username':'member','password':'wrong'})
        response=self.client.post('/login/',{'username':'member','password':'A-Fresh-Strong-Password!'})
        self.assertEqual(response.status_code,429)

    def test_security_headers(self):
        response=self.client.get('/')
        self.assertIn("frame-ancestors 'none'",response['Content-Security-Policy'])
        self.assertEqual(response['Cache-Control'],'no-store, private')


    def test_single_member_report_and_void_totals(self):
        payment = record_payment(self.payload, self.secretary)
        record_payment({**self.payload, 'member_id': self.other.pk, 'request_key': str(uuid4())}, self.secretary)
        url = f'/api/members/{self.member.pk}/report/'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        summary = dict(workbook['Summary'].values)
        self.assertEqual(summary['Total paid (excludes voids)'], 100)
        self.assertEqual(workbook['Payments'].max_row, 2)
        self.assertEqual(workbook['Months covered'].max_row, 5)
        self.assertNotIn('Other', str(list(workbook['Summary'].values)))
        self.assertEqual(workbook['Payments']['C2'].value, 100)
        self.assertEqual(workbook['Months covered']['C2'].value, 25)
        void_payment(payment.pk, self.secretary, 'Wrong payment recorded')
        workbook = load_workbook(io.BytesIO(self.client.get(url).content))
        self.assertEqual(dict(workbook['Summary'].values)['Total paid (excludes voids)'], 0)
        self.assertEqual(workbook['Payments']['F2'].value, 'VOID')
        self.assertTrue(AuditEvent.objects.filter(action='member.report_exported', entity_id=str(self.member.pk)).exists())

    def test_member_report_access_empty_and_formula_safety(self):
        url = f'/api/members/{self.member.pk}/report/'
        self.member.full_name = '=1+1'
        self.member.save()
        workbook = load_workbook(io.BytesIO(self.client.get(url).content))
        self.assertEqual(workbook['Summary']['B4'].data_type, 's')
        self.assertEqual(workbook['Payments'].max_row, 1)
        self.assertEqual(self.client.get('/api/members/999999/report/').status_code, 404)
        for user in [self.reader, self.auditor]:
            self.client.force_login(user)
            self.assertEqual(self.client.get(url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 401)
