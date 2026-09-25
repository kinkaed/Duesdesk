import csv
import io
import json
import hashlib
import logging
from datetime import date, timedelta
from decimal import Decimal
from functools import wraps
from uuid import uuid5, NAMESPACE_URL
from django.conf import settings
from django.contrib.auth import login
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.views import PasswordResetView
from django.core import signing
from django.core.exceptions import ValidationError
from django.db import transaction, IntegrityError, connection
from django.db.models import Sum
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods
from django.views.decorators.csrf import ensure_csrf_cookie
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .forms import SignupForm
from .models import Member, Payment, Allocation, DuesMonth, UserAccess, AuditEvent, Organisation, RecoveryAttempt, ImportBatch
from .access import role_for, visible_members, visible_payments
from .services import RATE, next_month, parse_month, parse_amount, plan_payment, record_payment, void_payment, audit

logger = logging.getLogger(__name__)

def api(view):
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return JsonResponse({'error': 'Your session ended. Please sign in again.'}, status=401)
        role = role_for(request.user)
        if not role:
            return JsonResponse({'error': 'Your account has no assigned access. Contact the secretary to assign your account role.'}, status=403)
        if request.method not in ('GET','HEAD') and role != 'secretary':
            return JsonResponse({'error': 'Your account is read-only.'}, status=403)
        try:
            return view(request, *args, **kwargs)
        except (ValueError, TypeError, ValidationError, Member.DoesNotExist, Payment.DoesNotExist) as error:
            message = '; '.join(error.messages) if isinstance(error, ValidationError) else str(error)
            return JsonResponse({'error': message or 'Check the entered values.'}, status=400)
        except IntegrityError:
            return JsonResponse({'error': 'A conflicting record already exists. Refresh and try again.'}, status=409)
    return wrapped

def secretary_only(request):
    return role_for(request.user) == 'secretary'

def body(request):
    try:
        data = json.loads(request.body)
        if not isinstance(data, dict): raise ValueError()
        return data
    except (ValueError, UnicodeDecodeError):
        raise ValueError('Send a valid form.')

def member_rows(month, user=None):
    people = list((visible_members(user) if user else Member.objects.all()).order_by('full_name'))
    amounts = Allocation.objects.filter(payment__voided_at__isnull=True, payment__member__in=people, dues_month__month__lte=month).values('payment__member_id','dues_month__month').annotate(total=Sum('amount'))
    paid = {(r['payment__member_id'], r['dues_month__month']): r['total'] for r in amounts}
    rates = dict(DuesMonth.objects.filter(month__lte=month).values_list('month','amount_due'))
    rows = []
    for person in people:
        start = person.joined.replace(day=1)
        applicable = start <= month and (not person.billing_end or month <= person.billing_end)
        amount = paid.get((person.pk, month), Decimal('0'))
        balance = max(Decimal('0'), rates.get(month,RATE) - amount) if applicable else Decimal('0')
        status = ('Paid' if balance == 0 else 'Partial' if amount else 'Unpaid') if applicable else 'Not due'
        arrears = Decimal('0')
        cursor = start
        last = min(month, person.billing_end) if person.billing_end else month
        while cursor <= last:
            arrears += max(Decimal('0'), rates.get(cursor,RATE) - paid.get((person.pk,cursor),Decimal('0')))
            cursor = next_month(cursor)
        rows.append({'id':person.pk,'code':person.code,'name':person.full_name,'phone':person.phone,'email':person.email,'joined':person.joined.isoformat(),'billing_end':person.billing_end.strftime('%Y-%m') if person.billing_end else '', 'member_status':person.status,'paid':str(amount),'balance':str(balance),'arrears':str(arrears),'status':status})
    return rows

def payment_json(payment):
    return {'id':payment.pk,'receipt':payment.receipt_number,'member_id':payment.member_id,'member':payment.member_name_snapshot or payment.member.full_name,'amount':str(payment.amount_received),'date':payment.payment_date.isoformat(),'method':payment.method,'reference':payment.reference,'notes':payment.notes,'voided':bool(payment.voided_at),'void_reason':payment.void_reason,'allocations':[{'month':a.dues_month.month.strftime('%B %Y'),'amount':str(a.amount)} for a in payment.allocations.all()]}

@require_http_methods(['GET', 'POST'])
def signup(request):
    if request.user.is_authenticated:
        return redirect(settings.LOGIN_REDIRECT_URL)
    form = SignupForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            user = form.save()
            UserAccess.objects.create(user=user, role='secretary')
        login(request, user, backend='django.contrib.auth.backends.ModelBackend')
        return redirect(settings.LOGIN_REDIRECT_URL)
    return render(request, 'registration/signup.html', {'form': form})


@login_required
def home(request):
    if not role_for(request.user): return HttpResponse('Your account has no assigned access. Contact the secretary to assign your account role.',status=403)
    return render(request,'react.html')

@api
@require_GET
@ensure_csrf_cookie
def session_info(request):
    org=Organisation.objects.first()
    return JsonResponse({'username':request.user.get_full_name() or request.user.username,'role':role_for(request.user),'user_id':request.user.pk,'today':timezone.localdate().isoformat(),'demo':settings.DEMO_MODE,'organisation':org.name if org else 'Membership Association'})

@api
@require_GET
def overview(request):
    month = parse_month(request.GET.get('month',timezone.localdate().strftime('%Y-%m')))
    rows = member_rows(month,request.user)
    due = [r for r in rows if r['status']!='Not due']
    records = visible_payments(request.user)
    transactions = records.filter(voided_at__isnull=True,payment_date__gte=month,payment_date__lt=next_month(month))
    collections = transactions.aggregate(total=Sum('amount_received'))['total'] or 0
    assigned = Allocation.objects.filter(payment__in=records,payment__voided_at__isnull=True,dues_month__month=month).aggregate(total=Sum('amount'))['total'] or 0
    recent = records.select_related('member').prefetch_related('allocations__dues_month').order_by('-created_at')[:8]
    return JsonResponse({'members':rows,'role':role_for(request.user),'metrics':{'collections':str(collections),'assigned':str(assigned),'outstanding':str(sum(Decimal(r['balance']) for r in due)),'arrears':str(sum(Decimal(r['arrears']) for r in rows)),'paid':sum(r['status']=='Paid' for r in due),'active':len(due)},'recent':[payment_json(p) for p in recent]})

def fill_member(person,data):
    try: joined=date.fromisoformat(data.get('joined',''))
    except (TypeError,ValueError): raise ValueError('Enter a joining date.')
    if not date(2000,1,1)<=joined<=timezone.localdate(): raise ValueError('Joining date must be between 2000 and today.')
    end = parse_month(data['billing_end']) if data.get('billing_end') else None
    if end and end<joined.replace(day=1): raise ValueError('Last billable month cannot be before the joining month.')
    if person.pk and person.payments.exists():
        if person.joined!=joined: raise ValueError('Joining date is locked after the first payment. Contact your administrator to correct historical billing.')
        if end and Allocation.objects.filter(payment__member=person,payment__voided_at__isnull=True,dues_month__month__gt=end).exists():
            raise ValueError('There are payments beyond the chosen last billable month. Correct those payments first.')
    person.full_name=str(data.get('name','')).strip()
    person.phone=str(data.get('phone','')).strip()
    person.email=str(data.get('email','')).strip()
    person.joined=joined
    person.status=data.get('status','Active')
    person.billing_end=end
    person.full_clean()
    return person

@api
@require_http_methods(['GET','POST'])
def members(request):
    if request.method=='GET':return JsonResponse({'members':member_rows(timezone.localdate().replace(day=1),request.user)})
    with transaction.atomic():
        person=fill_member(Member(),body(request));person.save()
        audit(request.user,'member.created',person,{'name':person.full_name,'joined':person.joined})
    return JsonResponse({'id':person.pk,'name':person.full_name,'code':person.code},status=201)

@api
@require_http_methods(['GET','POST'])
def member_detail(request,pk):
    if request.method=='POST':
        with transaction.atomic():
            member=get_object_or_404(visible_members(request.user).select_for_update(),pk=pk)
            before={'name':member.full_name,'phone':member.phone,'email':member.email,'status':member.status,'billing_end':member.billing_end}
            fill_member(member,body(request)).save()
            audit(request.user,'member.updated',member,{'before':before,'after':{'name':member.full_name,'phone':member.phone,'email':member.email,'status':member.status,'billing_end':member.billing_end}})
        return JsonResponse({'id':member.pk,'name':member.full_name})
    member=get_object_or_404(visible_members(request.user),pk=pk)
    records=member.payments.select_related('member').prefetch_related('allocations__dues_month').order_by('-payment_date','-id')
    total=records.filter(voided_at__isnull=True).aggregate(total=Sum('amount_received'))['total'] or 0
    row=next(r for r in member_rows(timezone.localdate().replace(day=1),request.user) if r['id']==member.pk)
    return JsonResponse({**row,'total':str(total),'payments':[payment_json(p) for p in records]})

@api
@require_http_methods(['POST'])
def payment_preview(request):
    data=body(request)
    member=Member.objects.get(pk=int(data.get('member_id',0)))
    plan=plan_payment(member,parse_amount(data.get('amount')),parse_month(data.get('start_month')))
    return JsonResponse({'allocations':[{'month':p['month'].strftime('%B %Y'),'amount':str(p['amount']),'status':p['status']} for p in plan]})

@api
@require_http_methods(['GET','POST'])
def payments(request):
    if request.method=='GET':
        page=max(1,int(request.GET.get('page',1)))
        items=visible_payments(request.user).select_related('member').prefetch_related('allocations__dues_month').order_by('-payment_date','-id')
        return JsonResponse({'payments':[payment_json(p) for p in items[(page-1)*100:page*100]],'page':page,'has_more':items.count()>page*100})
    item=record_payment(body(request),request.user)
    return JsonResponse({'id':item.pk,'receipt':item.receipt_number,'amount':str(item.amount_received)},status=201)

@api
@require_http_methods(['POST'])
def void(request,pk):
    item=void_payment(pk,request.user,body(request).get('reason',''))
    return JsonResponse({'id':item.pk,'receipt':item.receipt_number})

@login_required
def receipt(request,pk):
    item=get_object_or_404(visible_payments(request.user).select_related('member').prefetch_related('allocations__dues_month'),pk=pk)
    return render(request,'receipt.html',{'payment':item})

def safe_cell(value):
    value=str(value)
    return "'"+value if value.lstrip().startswith(('=','+','-','@')) else value

def report_data(request):
    month=parse_month(request.GET.get('month',timezone.localdate().strftime('%Y-%m')))
    kind=request.GET.get('kind','balances')
    if kind=='payments':
        header=['Receipt','Member ID','Member','Received','Method','Amount GHS','Reference','Status','Void reason']
        payments=visible_payments(request.user).filter(payment_date__gte=month,payment_date__lt=next_month(month)).select_related('member').order_by('id')
        rows=[[p.receipt_number,p.member.code,p.member_name_snapshot or p.member.full_name,p.payment_date.isoformat(),p.method,float(p.amount_received),p.reference,'VOID' if p.voided_at else 'Valid',p.void_reason] for p in payments]
    else:
        header=['Member ID','Name','Month','Paid GHS','Outstanding GHS','Arrears through month GHS','Payment status','Member status']
        rows=[[r['code'],r['name'],month.strftime('%Y-%m'),float(r['paid']),float(r['balance']),float(r['arrears']),r['status'],r['member_status']] for r in member_rows(month,request.user)]
    return month,header,[[safe_cell(v) if isinstance(v,str) else v for v in row] for row in rows]

@api
@require_GET
def export_csv(request):
    month,header,rows=report_data(request)
    response=HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition']=f'attachment; filename="dues-{month:%Y-%m}.csv"'
    response.write('\ufeff')
    writer=csv.writer(response);writer.writerow(header);writer.writerows(rows)
    audit(request.user,'report.exported',request.user,{'format':'csv','month':month,'rows':len(rows)})
    return response

@api
@require_GET
def export_excel(request):
    month,header,rows=report_data(request)
    book=Workbook();sheet=book.active;sheet.title='Monthly report'
    sheet.append(header)
    for row in rows:sheet.append(row)
    for cell in sheet[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='214F43')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width=min(45,max(18,max(len(str(c.value or '')) for c in column)+2))
    output=io.BytesIO();book.save(output)
    response=HttpResponse(output.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="dues-{month:%Y-%m}.xlsx"'
    audit(request.user,'report.exported',request.user,{'format':'xlsx','month':month,'rows':len(rows)})
    return response

@api
@require_http_methods(['GET','POST'])
def organisation_settings(request):
    if not secretary_only(request):return HttpResponse(status=403)
    item,_=Organisation.objects.get_or_create(pk=1)
    if request.method=='POST':
        with transaction.atomic():
            data=body(request)
            item.name=str(data.get('name','')).strip();item.contact=str(data.get('contact','')).strip();item.receipt_footer=str(data.get('receipt_footer','')).strip()
            item.full_clean();item.save();audit(request.user,'organisation.updated',item,{'name':item.name})
    return JsonResponse({'name':item.name,'contact':item.contact,'receipt_footer':item.receipt_footer})

@api
@require_GET
def audit_log(request):
    if role_for(request.user) not in ('secretary','auditor'):return HttpResponse(status=403)
    page=max(1,int(request.GET.get('page',1)))
    records=AuditEvent.objects.select_related('actor')
    return JsonResponse({'events':[{'id':e.pk,'date':e.created_at.isoformat(),'actor':e.actor.username if e.actor else 'System','action':e.action,'entity':e.entity,'entity_id':e.entity_id,'details':json.loads(e.details)} for e in records[(page-1)*100:page*100]],'page':page,'has_more':records.count()>page*100})

@api
@require_http_methods(['GET','POST'])
def accounts(request):
    if not secretary_only(request):return HttpResponse(status=403)
    if request.method=='GET':
        return JsonResponse({'users':[{'id':u.pk,'username':u.username,'email':u.email,'role':role_for(u) or 'None','active':u.is_active} for u in User.objects.select_related('access').order_by('username')]})
    data=body(request)
    with transaction.atomic():
        user=User(username=str(data.get('username','')).strip(),email=str(data.get('email','')).strip())
        if not user.email:raise ValueError('Email is required for account recovery.')
        if User.objects.filter(email__iexact=user.email).exists():raise ValueError('An account already uses this email.')
        role=data.get('role')
        if role not in ('secretary','auditor','member'):raise ValueError('Choose a valid role.')
        member=None
        if role=='member':member=Member.objects.get(pk=int(data.get('member_id',0)))
        password=data.get('password','');validate_password(password,user)
        user.set_password(password);user.full_clean();user.save()
        access=UserAccess(user=user,role=role,member=member);access.full_clean();access.save()
        audit(request.user,'account.created',user,{'role':role,'member_id':member.pk if member else None})
    return JsonResponse({'id':user.pk,'username':user.username},status=201)

@api
@require_http_methods(['POST'])
def disable_account(request,pk):
    if not secretary_only(request):return HttpResponse(status=403)
    if pk==request.user.pk:raise ValueError('You cannot disable your own account.')
    with transaction.atomic():
        user=get_object_or_404(User.objects.select_for_update(),pk=pk)
        if user.is_superuser and not request.user.is_superuser:raise ValueError('Only a superuser can disable a superuser.')
        user.is_active=False;user.save(update_fields=['is_active']);audit(request.user,'account.disabled',user)
    return JsonResponse({'ok':True})

@api
@require_GET
def import_template(request):
    if not secretary_only(request):return HttpResponse(status=403)
    kind=request.GET.get('kind','members')
    text='name,phone,email,joined,status,billing_end\nExample Member,0240000000,,2026-09-01,Active,\n' if kind=='members' else 'member_id,amount,start_month,payment_date,method,reference,notes\n1,100,2026-09,2026-09-22,Cash,,Example import\n'
    response=HttpResponse(text,content_type='text/csv');response['Content-Disposition']=f'attachment; filename="{kind}-template.csv"';return response

@api
@require_http_methods(['POST'])
def import_preview(request):
    file=request.FILES.get('file');kind=request.POST.get('kind')
    if kind not in ('members','payments'):raise ValueError('Choose members or payments.')
    if not file or file.size>1048576:raise ValueError('Choose a CSV file smaller than 1 MB.')
    try:
        reader=csv.DictReader(io.StringIO(file.read().decode('utf-8-sig')))
        required={'name','joined'} if kind=='members' else {'member_id','amount','start_month','payment_date','method'}
        if not required.issubset(reader.fieldnames or []):raise ValueError('Missing required column headers. Use the template.')
        rows=list(reader)
    except (UnicodeDecodeError,csv.Error):raise ValueError('Use a valid UTF-8 CSV file.')
    if not 1<=len(rows)<=500:raise ValueError('Import between 1 and 500 rows at a time.')
    for i,row in enumerate(rows,2):
        if None in row or any(v is None for v in row.values()):raise ValueError(f'Row {i}: column count does not match the header.')
        try:
            if kind=='members':fill_member(Member(),row)
            else:
                member=Member.objects.get(pk=int(row['member_id']))
                plan_payment(member,parse_amount(row['amount']),parse_month(row['start_month']))
                payment_date=date.fromisoformat(row['payment_date'])
                if not date(2000,1,1)<=payment_date<=timezone.localdate():raise ValueError('Invalid payment date.')
                if row['method'] not in dict(Payment._meta.get_field('method').choices):raise ValueError('Invalid method.')
                for field in ('reference', 'notes'):
                    limit = Payment._meta.get_field(field).max_length
                    if len(str(row.get(field, '')).strip()) > limit:
                        raise ValueError(f'{field.capitalize()} must be {limit} characters or fewer.')
        except (ValueError,ValidationError,Member.DoesNotExist) as e:raise ValueError(f'Row {i}: {e}')
    token=signing.dumps({'rows':rows,'kind':kind,'user':request.user.pk},salt='csv-import',compress=True)
    return JsonResponse({'token':token,'count':len(rows),'preview':rows[:8],'kind':kind})

@api
@require_http_methods(['POST'])
def import_commit(request):
    try:payload=signing.loads(body(request).get('token',''),salt='csv-import',max_age=600)
    except signing.BadSignature:raise ValueError('The preview expired. Upload the file again.')
    if payload['user']!=request.user.pk:raise ValueError('This preview belongs to another user.')
    digest=hashlib.sha256(json.dumps({'kind':payload['kind'],'rows':payload['rows']},sort_keys=True).encode()).hexdigest()
    with transaction.atomic():
        if ImportBatch.objects.filter(digest=digest).exists():raise ValueError('This exact file has already been imported.')
        batch=ImportBatch.objects.create(digest=digest,kind=payload['kind'],row_count=len(payload['rows']),created_by=request.user)
        for row_index,row in enumerate(payload['rows'],1):
            if payload['kind']=='members':
                member=fill_member(Member(),row);member.save();audit(request.user,'member.imported',member,{'batch':batch.pk})
            else:
                row['request_key']=str(uuid5(NAMESPACE_URL,f'duesdesk:{digest}:{row_index}'))
                record_payment(row,request.user)
        audit(request.user,'import.completed',batch,{'kind':batch.kind,'rows':batch.row_count})
    return JsonResponse({'count':batch.row_count})

class RecoveryView(PasswordResetView):
    template_name='registration/account_form.html'
    email_template_name='registration/password_reset_email.txt'
    subject_template_name='registration/password_reset_subject.txt'
    extra_context={'heading':'Reset your password','description':'Enter the email address attached to your account.','button':'Send reset link'}
    def form_valid(self,form):
        # Do not reveal whether an email is registered. Limit repeated emails to one per 5 minutes.
        email=form.cleaned_data['email'].strip().lower()
        key=hashlib.sha256(email.encode()).hexdigest()
        with transaction.atomic():
            item,_=RecoveryAttempt.objects.select_for_update().get_or_create(key=key,defaults={'requested_at':timezone.now()-timedelta(days=1)})
            if item.requested_at>timezone.now()-timedelta(minutes=5):return redirect('password_reset_done')
            item.requested_at=timezone.now();item.save()
        try:return super().form_valid(form)
        except Exception:
            logger.exception('Password recovery delivery failed')
            return redirect('password_reset_done')

@require_GET
def health(request):
    try:
        with connection.cursor() as cursor:cursor.execute('SELECT 1');cursor.fetchone()
        return JsonResponse({'status':'ok'})
    except Exception:return JsonResponse({'status':'unavailable'},status=503)


@api
@require_GET
def member_report(request, pk):
    """Secretary-requested statement: one receipt row, separate monthly allocations."""
    if not secretary_only(request):
        return HttpResponse(status=403)
    member = get_object_or_404(Member, pk=pk)
    today = timezone.localdate()
    current_month = today.replace(day=1)
    payments = list(member.payments.prefetch_related('allocations__dues_month').order_by('payment_date', 'pk'))
    allocated = {}
    for payment in payments:
        if not payment.voided_at:
            for allocation in payment.allocations.all():
                month = allocation.dues_month.month
                allocated[month] = allocated.get(month, Decimal('0')) + allocation.amount
    rates = dict(DuesMonth.objects.values_list('month', 'amount_due'))
    monthly_rows = []
    arrears = Decimal('0')
    cursor = member.joined.replace(day=1)
    last_due = min(current_month, member.billing_end) if member.billing_end else current_month
    last = max([last_due, *allocated.keys()])
    while cursor <= last:
        due = rates.get(cursor, RATE) if cursor <= last_due else Decimal('0')
        paid = allocated.get(cursor, Decimal('0'))
        outstanding = max(due - paid, Decimal('0'))
        arrears += outstanding
        monthly_rows.append([cursor.strftime('%B %Y'), due, paid, outstanding,
                             'Advance payment' if cursor > last_due else 'Paid' if not outstanding else 'Partial' if paid else 'Unpaid'])
        cursor = next_month(cursor)
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    org = Organisation.objects.first()
    for row in [
        ['Member payment report', 'Value'],
        ['Organisation', org.name if org else 'Membership Association'],
        ['Member ID', member.code], ['Member name', member.full_name],
        ['Generated on', today.isoformat()], ['Currency', 'GHS'],
        ['Total paid (excludes voids)', sum((p.amount_received for p in payments if not p.voided_at), Decimal('0'))],
        ['Outstanding through ' + current_month.strftime('%B %Y'), arrears],
        ['Report coverage', 'All recorded payments; outstanding dues through the current month'],
        ['Note', 'Voided payments are retained for reference and excluded from totals.'],
    ]:
        summary.append([(safe_cell(value) if isinstance(value, str) else value) for value in row])
    ledger = workbook.create_sheet('Payments')
    ledger.append(['Receipt', 'Payment date', 'Amount (GHS)', 'Method', 'Reference', 'Status', 'Notes', 'Void reason'])
    allocations = workbook.create_sheet('Months covered')
    allocations.append(['Receipt', 'Covered month', 'Allocated amount (GHS)', 'Receipt status'])
    for payment in payments:
        status = 'VOID' if payment.voided_at else 'Valid'
        ledger.append([(safe_cell(value) if isinstance(value, str) else value) for value in [payment.receipt_number, payment.payment_date.isoformat(),
            payment.amount_received, payment.method, payment.reference, status, payment.notes, payment.void_reason]])
        for allocation in sorted(payment.allocations.all(), key=lambda a: a.dues_month.month):
            allocations.append([payment.receipt_number, allocation.dues_month.month.strftime('%B %Y'), allocation.amount, status])
    balances = workbook.create_sheet('Monthly balances')
    balances.append(['Month', 'Due to date (GHS)', 'Paid towards month (GHS)', 'Outstanding (GHS)', 'Status'])
    for row in monthly_rows:
        balances.append(row)
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='166534')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(70, max(18, max(len(str(c.value or '')) for c in column) + 2))
            for cell in column[1:]:
                if isinstance(cell.value, Decimal):
                    cell.number_format = '#,##0.00'
    output = io.BytesIO()
    workbook.save(output)
    audit(request.user, 'member.report_exported', member, {'format': 'xlsx', 'payments': len(payments)})
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{member.code}-payment-report-{today.isoformat()}.xlsx"'
    return response
